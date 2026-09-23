import csv
from decimal import Decimal
from io import BytesIO, StringIO
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from attendance_api.models import (
    AuditLog,
    ClassGroup,
    Course,
    Enrollment,
    ScoreItem,
    ScoreRecord,
    ScoreSettings,
    Student,
    User,
)
from attendance_api.security.passwords import hash_password


@pytest.fixture
def scores(client: TestClient, db_session: Session) -> tuple[TestClient, str, list[str]]:
    teacher = User(
        username=f"scores-{uuid4()}",
        password_hash=hash_password("teacher-password-123"),
        display_name="成绩教师",
        role="TEACHER",
        status="ACTIVE",
        must_change_password=False,
    )
    db_session.add(teacher)
    db_session.flush()
    course = Course(
        owner_teacher_id=teacher.id, name="程序设计", code="CS", term="2026秋", status="ACTIVE"
    )
    db_session.add(course)
    db_session.flush()
    group = ClassGroup(course_id=course.id, name="一班", status="ACTIVE")
    db_session.add(group)
    db_session.flush()
    ids = []
    for index in range(2):
        student = Student(student_number=f"score-{uuid4()}", name=f"学生{index}")
        db_session.add(student)
        db_session.flush()
        enrollment = Enrollment(class_group_id=group.id, student_id=student.id, status="ACTIVE")
        db_session.add(enrollment)
        db_session.flush()
        ids.append(enrollment.id)
    db_session.commit()
    response = client.post(
        "/api/v1/auth/login",
        json={
            "username": teacher.username,
            "password": "teacher-password-123",
        },
    )
    assert response.status_code == 200
    client.headers["Authorization"] = f"Bearer {response.json()['access_token']}"
    return client, f"/api/v1/classes/{group.id}/scores", ids


def add_item(
    client: TestClient,
    url: str,
    version: int = 0,
    category: str = "HOMEWORK",
    default_points: str | None = None,
) -> dict:
    payload = {
        "expected_version": version,
        "category": category,
        "name": " 作业1 ",
        "occurred_on": "2026-09-22",
        "description": "说明",
    }
    if default_points is not None:
        payload["default_points"] = default_points
    response = client.post(
        url + "/items",
        json=payload,
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_project_default_points_seed_records_and_remain_editable(scores) -> None:
    client, url, ids = scores
    book = add_item(client, url, default_points="3")
    item = book["items"][0]

    assert item["default_points"] == "3"
    assert [record["points"] for record in book["records"]] == ["3", "3"]

    saved = client.put(
        f"{url}/items/{item['id']}/records",
        json={
            "expected_version": 1,
            "records": [{"enrollment_id": ids[0], "points": "5"}],
        },
    )
    assert saved.status_code == 200, saved.text
    saved_record = next(
        record for record in saved.json()["records"] if record["enrollment_id"] == ids[0]
    )
    assert saved_record["points"] == "5"

    changed = client.patch(
        f"{url}/items/{item['id']}",
        json={"expected_version": 2, "default_points": "4"},
    )
    assert changed.status_code == 200, changed.text
    assert changed.json()["items"][0]["default_points"] == "4"
    records = {record["enrollment_id"]: record["points"] for record in changed.json()["records"]}
    assert records == {ids[0]: "5", ids[1]: "3"}


def set_rules(client: TestClient, url: str, version: int, base: str | None = "70") -> dict:
    response = client.put(
        url + "/settings",
        json={
            "expected_version": version,
            "base_score": base,
            "factors": {"HOMEWORK": "2", "CLASSROOM": "1", "LAB": "0.5", "OTHER": "0"},
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_new_book_is_read_only_query_and_defaults_are_unconfigured(scores, db_session) -> None:
    client, url, ids = scores
    response = client.get(url)
    assert response.status_code == 200
    book = response.json()
    assert book["version"] == 0
    assert book["settings"]["base_score"] == "70"
    assert not book["rules_ready"]
    assert book["settings"]["factors"] == dict.fromkeys(["HOMEWORK", "CLASSROOM", "LAB", "OTHER"])
    assert {student["enrollment_id"] for student in book["students"]} == set(ids)
    assert (
        db_session.scalar(
            select(func.count())
            .select_from(ScoreSettings)
            .where(ScoreSettings.class_group_id == book["class_group"]["id"])
        )
        == 0
    )


def test_batch_decimal_points_settings_recalculation_and_null_zero(scores) -> None:
    client, url, ids = scores
    book = add_item(client, url)
    item_id = book["items"][0]["id"]
    assert book["items"][0]["name"] == "作业1"
    assert len(book["records"]) == 2
    assert all(record["points"] is None for record in book["records"])
    set_rules(client, url, 1)
    response = client.put(
        f"{url}/items/{item_id}/records",
        json={
            "expected_version": 2,
            "records": [
                {"enrollment_id": ids[0], "points": "0.5", "note": "部分完成"},
                {"enrollment_id": ids[1], "points": 0, "note": "未完成"},
            ],
        },
    )
    assert response.status_code == 200, response.text
    book = response.json()
    summaries = {row["enrollment_id"]: row for row in book["summaries"]}
    assert summaries[ids[0]]["final_score"] == "71.00"
    assert summaries[ids[1]]["final_score"] == "70.00"
    assert book["version"] == 3
    assert all(isinstance(record["points"], str) for record in book["records"])
    book = set_rules(client, url, 3, "80")
    assert sorted(row["final_score"] for row in book["summaries"]) == ["80.00", "81.00"]
    cleared = client.put(
        f"{url}/items/{item_id}/records",
        json={
            "expected_version": 4,
            "records": [
                {"enrollment_id": ids[0], "points": None, "note": "待复评"},
            ],
        },
    )
    assert cleared.status_code == 200
    row = next(row for row in cleared.json()["summaries"] if row["enrollment_id"] == ids[0])
    assert row["status"] == "RECORDS_PENDING"
    assert row["final_score"] == "80.00"


def test_version_conflict_and_invalid_batch_are_atomic(scores, db_session) -> None:
    client, url, ids = scores
    book = add_item(client, url)
    item_id = book["items"][0]["id"]
    for version, records, status in [
        (0, [{"enrollment_id": ids[0], "points": "2"}], 409),
        (
            1,
            [
                {"enrollment_id": ids[0], "points": "2"},
                {"enrollment_id": str(uuid4()), "points": "1"},
            ],
            404,
        ),
        (
            1,
            [{"enrollment_id": ids[0], "points": "2"}, {"enrollment_id": ids[0], "points": "1"}],
            400,
        ),
    ]:
        response = client.put(
            f"{url}/items/{item_id}/records",
            json={
                "expected_version": version,
                "records": records,
            },
        )
        assert response.status_code == status, response.text
        current = client.get(url).json()
        assert current["version"] == 1
        assert all(record["points"] is None for record in current["records"])
    assert (
        db_session.scalar(
            select(func.count())
            .select_from(AuditLog)
            .where(
                AuditLog.action == "SCORE_RECORD_UPDATED",
                AuditLog.entity_id == book["class_group"]["id"],
            )
        )
        == 0
    )


def test_removed_students_history_and_restore(scores, db_session) -> None:
    client, url, ids = scores
    book = add_item(client, url)
    first_item = book["items"][0]["id"]
    set_rules(client, url, 1)
    saved = client.put(
        f"{url}/items/{first_item}/records",
        json={
            "expected_version": 2,
            "records": [
                {"enrollment_id": ids[0], "points": "1.5", "note": "完成"},
            ],
        },
    )
    assert saved.status_code == 200
    assert (
        client.patch(f"/api/v1/enrollments/{ids[0]}/status", json={"status": "REMOVED"}).status_code
        == 200
    )
    book = add_item(client, url, 3, "LAB")
    assert len(book["records"]) == 3
    old = next(row for row in book["summaries"] if row["enrollment_id"] == ids[0])
    assert old["final_score"] == "73.00"
    rejected = client.put(
        f"{url}/items/{first_item}/records",
        json={
            "expected_version": 4,
            "records": [
                {"enrollment_id": ids[0], "points": "2"},
            ],
        },
    )
    assert rejected.status_code == 409
    history = client.get(url + "/history", params={"enrollment_id": ids[0]}).json()["items"]
    assert {row["action"] for row in history} == {
        "SCORE_ITEM_CREATED",
        "SCORE_SETTINGS_UPDATED",
        "SCORE_RECORD_UPDATED",
    }
    changed = next(row for row in history if row["action"] == "SCORE_RECORD_UPDATED")
    assert changed["before_value"]["points"] is None
    assert changed["after_value"]["points"] == "1.5"
    assert changed["actor_name"] == "成绩教师"
    assert (
        client.patch(f"/api/v1/enrollments/{ids[0]}/status", json={"status": "ACTIVE"}).status_code
        == 200
    )
    restored = client.get(url).json()
    old = next(row for row in restored["summaries"] if row["enrollment_id"] == ids[0])
    assert old["status"] == "RECORDS_PENDING"
    saved = client.put(
        f"{url}/items/{first_item}/records",
        json={
            "expected_version": 4,
            "records": [
                {"enrollment_id": ids[0], "points": "2"},
            ],
        },
    )
    assert saved.status_code == 200


@pytest.mark.parametrize("scope", ["class_group", "course"])
def test_archived_scope_is_readonly_but_exportable(scores, scope) -> None:
    client, url, _ = scores
    book = add_item(client, url)
    resource = "classes" if scope == "class_group" else "courses"
    assert (
        client.patch(
            f"/api/v1/{resource}/{book[scope]['id']}",
            json={
                "status": "ARCHIVED",
            },
        ).status_code
        == 200
    )
    assert client.get(url).json()["readonly"]
    response = client.put(
        url + "/settings",
        json={
            "expected_version": 1,
            "base_score": "70",
            "factors": {c: "1" for c in ("HOMEWORK", "CLASSROOM", "LAB", "OTHER")},
        },
    )
    assert response.status_code == 409
    assert client.get(url + "/export?format=csv&kind=details").status_code == 200


def test_item_update_and_cross_class_isolation(scores, db_session) -> None:
    client, url, ids = scores
    book = add_item(client, url)
    item_id = book["items"][0]["id"]
    other = client.post(
        f"/api/v1/courses/{book['course']['id']}/classes", json={"name": "二班"}
    ).json()
    other_url = f"/api/v1/classes/{other['id']}/scores"
    assert client.get(other_url).json()["version"] == 0
    response = client.patch(
        f"{other_url}/items/{item_id}", json={"expected_version": 0, "name": "越界"}
    )
    assert response.status_code == 404
    other_book = add_item(client, other_url)
    other_item = other_book["items"][0]["id"]
    response = client.put(
        f"{other_url}/items/{other_item}/records",
        json={
            "expected_version": 1,
            "records": [{"enrollment_id": ids[0], "points": "1"}],
        },
    )
    assert response.status_code == 404
    response = client.patch(
        f"{url}/items/{item_id}",
        json={
            "expected_version": 1,
            "name": "作业二",
            "occurred_on": "2026-09-23",
            "description": "新说明",
        },
    )
    assert response.status_code == 200
    assert response.json()["items"][0]["name"] == "作业二"
    assert client.get(other_url + "/history", params={"enrollment_id": ids[0]}).status_code == 404


def test_other_teacher_and_admin_cannot_access_scores(scores, db_session) -> None:
    client, url, _ = scores
    for role, status in [("TEACHER", 404), ("ADMIN", 403)]:
        user = User(
            username=f"foreign-{uuid4()}",
            password_hash=hash_password("other-password-123"),
            display_name="其他用户",
            role=role,
            status="ACTIVE",
            must_change_password=False,
        )
        db_session.add(user)
        db_session.commit()
        token = client.post(
            "/api/v1/auth/login",
            json={
                "username": user.username,
                "password": "other-password-123",
            },
        ).json()["access_token"]
        client.headers["Authorization"] = f"Bearer {token}"
        for suffix in ["", "/history", "/export?format=csv&kind=details"]:
            assert client.get(url + suffix).status_code == status


def test_export_pending_rows_and_formula_injection(scores, db_session) -> None:
    client, url, ids = scores
    book = add_item(client, url)
    item_id = book["items"][0]["id"]
    assert client.get(url + "/export?format=csv&kind=summary").status_code == 409
    response = client.put(
        f"{url}/items/{item_id}/records",
        json={
            "expected_version": 1,
            "records": [
                {"enrollment_id": ids[0], "points": "-0.5", "note": '\t=HYPERLINK("bad")'},
            ],
        },
    )
    assert response.status_code == 200
    details = client.get(url + "/export?format=csv&kind=details")
    rows = list(csv.reader(StringIO(details.content.decode("utf-8-sig"))))
    assert any('\'\t=HYPERLINK("bad")' in row for row in rows)
    assert any("尚未评价" in row and row[8] == "" for row in rows[1:])
    xlsx = client.get(url + "/export?format=xlsx&kind=details")
    workbook = load_workbook(BytesIO(xlsx.content))
    assert all(cell.data_type != "f" for row in workbook.active for cell in row)
    workbook.close()
    set_rules(client, url, 2)
    summary = client.get(url + "/export?format=csv&kind=summary")
    rows = list(csv.reader(StringIO(summary.content.decode("utf-8-sig"))))
    assert any(row[-1] == "待评完整" and row[-2] == "70.00" for row in rows[1:])
    assert any(row[-2] == "69.00" for row in rows[1:])


def test_explicit_class_deletion_cleans_scores_but_preserves_student(scores, db_session) -> None:
    client, url, _ = scores
    book = add_item(client, url)
    class_id = book["class_group"]["id"]
    student_id = book["students"][0]["student_id"]
    assert client.delete(f"/api/v1/classes/{class_id}").status_code == 409
    response = client.delete(f"/api/v1/classes/{class_id}?delete_related_data=true")
    assert response.status_code == 204, response.text
    for model in [ScoreSettings, ScoreItem]:
        assert (
            db_session.scalar(
                select(func.count()).select_from(model).where(model.class_group_id == class_id)
            )
            == 0
        )
    assert (
        db_session.scalar(
            select(func.count())
            .select_from(ScoreRecord)
            .where(ScoreRecord.item_id == book["items"][0]["id"])
        )
        == 0
    )
    assert db_session.get(Student, student_id) is not None


def test_decimal_validation_and_empty_settings(scores) -> None:
    client, url, ids = scores
    book = add_item(client, url)
    item_id = book["items"][0]["id"]
    for point in ["NaN", "Infinity", "0.00001", "10000000000000000"]:
        response = client.put(
            f"{url}/items/{item_id}/records",
            json={
                "expected_version": 1,
                "records": [{"enrollment_id": ids[0], "points": point}],
            },
        )
        assert response.status_code == 400
    result = set_rules(client, url, 1, None)
    assert result["settings"]["base_score"] is None
    assert not result["rules_ready"]
    assert all(row["final_score"] is None for row in result["summaries"])


def test_raw_export_retains_four_decimal_places(scores) -> None:
    client, url, ids = scores
    book = add_item(client, url)
    item_id = book["items"][0]["id"]
    response = client.put(
        f"{url}/items/{item_id}/records",
        json={
            "expected_version": 1,
            "records": [{"enrollment_id": ids[0], "points": "0.0001"}],
        },
    )
    assert response.status_code == 200
    summary = next(row for row in response.json()["summaries"] if row["enrollment_id"] == ids[0])
    assert Decimal(summary["categories"][0]["points"]) == Decimal("0.0001")
    details = client.get(url + "/export?format=csv&kind=details")
    rows = list(csv.reader(StringIO(details.content.decode("utf-8-sig"))))
    assert any(row[8] == "0.0001" for row in rows[1:])
    workbook = load_workbook(BytesIO(client.get(url + "/export?format=xlsx&kind=details").content))
    assert any(row[8].value == "0.0001" for row in list(workbook.active)[1:])
    workbook.close()
