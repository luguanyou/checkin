from dataclasses import dataclass
from datetime import date
from io import BytesIO
from urllib.parse import unquote

import pytest
from argon2 import PasswordHasher
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from attendance_api.models import (
    AttendanceRecord,
    AttendanceSession,
    ClassGroup,
    Course,
    Student,
    User,
)
from attendance_api.models.base import utc_now

PASSWORD = "teacher-password-123"
EXPORT_COLUMNS = [
    "课程",
    "课程代码",
    "班级",
    "日期",
    "姓名",
    "学号",
    "考勤状态",
    "点名时间",
    "最后修改时间",
    "最后修改人",
]


@dataclass
class ExportContext:
    session: AttendanceSession
    foreign_session: AttendanceSession
    draft_session: AttendanceSession
    client: TestClient


@pytest.fixture
def export_context(client: TestClient, db_session: Session) -> ExportContext:
    teacher = User(
        username="export-teacher",
        password_hash=PasswordHasher().hash(PASSWORD),
        display_name="导出教师",
        role="TEACHER",
        status="ACTIVE",
        must_change_password=False,
    )
    other = User(
        username="export-other",
        password_hash=PasswordHasher().hash(PASSWORD),
        display_name="其他教师",
        role="TEACHER",
        status="ACTIVE",
        must_change_password=False,
    )
    db_session.add_all([teacher, other])
    db_session.flush()

    def add_session(owner: User, code: str, status: str) -> AttendanceSession:
        course = Course(
            owner_teacher_id=owner.id,
            name="设计/课" if owner.id == teacher.id else "其他课",
            code=code,
            term="2026 秋季学期",
            status="ACTIVE",
        )
        db_session.add(course)
        db_session.flush()
        class_group = ClassGroup(course_id=course.id, name="一:班", status="ACTIVE")
        db_session.add(class_group)
        db_session.flush()
        session = AttendanceSession(
            class_group_id=class_group.id,
            session_date=date(2026, 9, 3 if code != "DRAFT" else 4),
            status=status,
            started_at=utc_now(),
            completed_at=utc_now() if status == "COMPLETED" else None,
            created_by=owner.id,
        )
        db_session.add(session)
        db_session.flush()
        student = Student(student_number=f"00{code}", name="原姓名")
        db_session.add(student)
        db_session.flush()
        db_session.add(
            AttendanceRecord(
                session_id=session.id,
                student_id=student.id,
                student_number_snapshot=f"00{code}",
                student_name_snapshot='=HYPERLINK("https://example.test")',
                class_name_snapshot="一:班",
                status="present",
                marked_at=utc_now(),
                last_modified_by=owner.id,
                last_modified_at=utc_now(),
                version=2,
            )
        )
        return session

    completed = add_session(teacher, "EXPORT", "COMPLETED")
    draft = add_session(teacher, "DRAFT", "DRAFT")
    foreign = add_session(other, "FOREIGN", "COMPLETED")
    db_session.commit()
    login = client.post(
        "/api/v1/auth/login", json={"username": teacher.username, "password": PASSWORD}
    )
    client.headers.update({"Authorization": f"Bearer {login.json()['access_token']}"})
    return ExportContext(completed, foreign, draft, client)


def test_csv_has_bom_fixed_columns_snapshot_values_and_safe_formulas(
    export_context: ExportContext,
) -> None:
    response = export_context.client.get(
        f"/api/v1/attendance-sessions/{export_context.session.id}/export",
        params={"format": "csv"},
    )

    assert response.status_code == 200
    assert response.content.startswith(b"\xef\xbb\xbf")
    text = response.content.decode("utf-8-sig")
    assert text.splitlines()[0].split(",") == EXPORT_COLUMNS
    assert "'=HYPERLINK" in text
    assert "00EXPORT" in text
    assert response.headers["X-Content-Type-Options"] == "nosniff"


def test_xlsx_student_numbers_are_string_cells(export_context: ExportContext) -> None:
    response = export_context.client.get(
        f"/api/v1/attendance-sessions/{export_context.session.id}/export",
        params={"format": "xlsx"},
    )

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=False, data_only=False)
    sheet = workbook.active
    assert sheet is not None
    assert [cell.value for cell in sheet[1]] == EXPORT_COLUMNS
    assert sheet.cell(row=2, column=6).value == "00EXPORT"
    assert sheet.cell(row=2, column=6).data_type == "s"
    assert sheet.cell(row=2, column=5).data_type == "s"
    workbook.close()


def test_download_filename_is_utf8_and_sanitized(export_context: ExportContext) -> None:
    response = export_context.client.get(
        f"/api/v1/attendance-sessions/{export_context.session.id}/export",
        params={"format": "csv"},
    )

    disposition = response.headers["Content-Disposition"]
    encoded = disposition.split("filename*=UTF-8''", 1)[1]
    filename = unquote(encoded)
    assert filename == "设计_课_一_班_2026-09-03_单次考勤.csv"
    assert "/" not in filename and ":" not in filename


def test_draft_unknown_format_and_foreign_session_are_rejected(
    export_context: ExportContext,
) -> None:
    draft = export_context.client.get(
        f"/api/v1/attendance-sessions/{export_context.draft_session.id}/export",
        params={"format": "csv"},
    )
    unknown = export_context.client.get(
        f"/api/v1/attendance-sessions/{export_context.session.id}/export",
        params={"format": "pdf"},
    )
    foreign = export_context.client.get(
        f"/api/v1/attendance-sessions/{export_context.foreign_session.id}/export",
        params={"format": "csv"},
    )

    assert draft.status_code == 409
    assert draft.json()["code"] == "ATTENDANCE_SESSION_NOT_COMPLETED"
    assert unknown.status_code == 400
    assert unknown.json()["code"] == "INVALID_REQUEST"
    assert foreign.status_code == 404
