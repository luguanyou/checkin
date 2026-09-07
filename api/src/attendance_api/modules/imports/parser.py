import csv
from dataclasses import dataclass
from io import StringIO
from pathlib import Path, PurePath
from typing import TypedDict
from zipfile import BadZipFile, ZipFile

import xlrd
from openpyxl import load_workbook

from attendance_api.errors import ApiError

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_ROWS = 500
MAX_UNCOMPRESSED_XLSX_SIZE = 50 * 1024 * 1024
XLS_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")
HEADER_ALIASES = {
    "学号": "student_number",
    "学生学号": "student_number",
    "student_number": "student_number",
    "studentnumber": "student_number",
    "姓名": "name",
    "学生姓名": "name",
    "name": "name",
    "性别": "gender",
    "gender": "gender",
    "班级": "class_name",
    "班级名称": "class_name",
    "class": "class_name",
    "class_name": "class_name",
    "专业": "major",
    "major": "major",
}
REQUIRED_HEADERS = {"student_number", "name", "class_name"}
FIELD_LIMITS = {
    "student_number": 50,
    "name": 100,
    "gender": 20,
    "class_name": 120,
    "major": 120,
}


@dataclass(frozen=True)
class ParsedRoster:
    rows: list[dict[str, str | int | None]]
    validation: "ValidationResult"


class ValidationResult(TypedDict):
    summary: dict[str, int]
    rows: list[dict[str, object]]


def sanitize_filename(filename: str) -> str:
    sanitized = PurePath(filename.replace("\\", "/")).name.strip()
    return sanitized[:255] or "upload"


def _parse_error() -> ApiError:
    return ApiError(400, "ROSTER_PARSE_FAILED", "名单文件无法解析")


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_csv(path: Path) -> list[list[object]]:
    data = path.read_bytes()
    if data.startswith(b"PK\x03\x04") or data.startswith(XLS_SIGNATURE) or b"\x00" in data:
        raise _parse_error()
    decoded: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise _parse_error()
    try:
        return [list(row) for row in csv.reader(StringIO(decoded, newline=""))]
    except csv.Error as exc:
        raise _parse_error() from exc


def _validate_xlsx_archive(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            total = sum(item.file_size for item in archive.infolist())
            compressed = sum(item.compress_size for item in archive.infolist())
            if total > MAX_UNCOMPRESSED_XLSX_SIZE:
                raise _parse_error()
            if compressed and total > 1024 * 1024 and total / compressed > 100:
                raise _parse_error()
    except BadZipFile as exc:
        raise _parse_error() from exc


def _read_xlsx(path: Path) -> list[list[object]]:
    if not path.read_bytes()[:4] == b"PK\x03\x04":
        raise _parse_error()
    _validate_xlsx_archive(path)
    try:
        with path.open("rb") as source:
            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                if sheet is None:
                    raise _parse_error()
                return [list(row) for row in sheet.iter_rows(values_only=True)]
            finally:
                workbook.close()
    except Exception as exc:
        if isinstance(exc, ApiError):
            raise
        raise _parse_error() from exc


def _read_xls(path: Path) -> list[list[object]]:
    if not path.read_bytes()[:8] == XLS_SIGNATURE:
        raise _parse_error()
    try:
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        sheet = workbook.sheet_by_index(0)
        return [list(sheet.row_values(index)) for index in range(sheet.nrows)]
    except Exception as exc:
        raise _parse_error() from exc


def _normalize_header(value: object) -> str | None:
    text = _cell_text(value).lower().replace(" ", "_")
    return HEADER_ALIASES.get(text)


def _validation_row(
    row: dict[str, str | int | None], target_class_name: str, seen: set[str]
) -> dict[str, object]:
    required_missing = [
        field for field in REQUIRED_HEADERS if not str(row.get(field) or "").strip()
    ]
    invalid_fields = [
        field for field, maximum in FIELD_LIMITS.items() if len(str(row.get(field) or "")) > maximum
    ]
    status = "valid"
    code: str | None = None
    fields: list[str] = []
    message: str | None = None
    if required_missing:
        status, code, fields, message = (
            "error",
            "MISSING_REQUIRED_FIELD",
            sorted(required_missing),
            "缺少必填字段",
        )
    elif invalid_fields:
        status, code, fields, message = (
            "error",
            "INVALID_FIELD",
            sorted(invalid_fields),
            "字段长度超出限制",
        )
    elif row["class_name"] != target_class_name:
        status, code, fields, message = (
            "error",
            "CLASS_MISMATCH",
            ["class_name"],
            "班级名称与目标班级不一致",
        )
    elif str(row["student_number"]) in seen:
        status, code, fields, message = (
            "duplicate",
            "DUPLICATE_IN_FILE",
            ["student_number"],
            "文件中存在重复学号",
        )
    if row["student_number"]:
        seen.add(str(row["student_number"]))
    return {**row, "status": status, "code": code, "fields": fields, "message": message}


def _summarize(rows: list[dict[str, object]]) -> dict[str, int]:
    return {
        "total": len(rows),
        "valid": sum(row["status"] == "valid" for row in rows),
        "errors": sum(row["status"] in {"error", "identity_conflict"} for row in rows),
        "duplicates": sum(row["status"] == "duplicate" for row in rows),
        "importable": sum(row["status"] == "valid" for row in rows),
    }


def parse_roster(path: Path, filename: str, target_class_name: str) -> ParsedRoster:
    if path.stat().st_size > MAX_FILE_SIZE:
        raise ApiError(400, "FILE_TOO_LARGE", "名单文件不能超过 10 MiB")
    extension = Path(sanitize_filename(filename)).suffix.lower()
    if extension not in {".csv", ".xlsx", ".xls"}:
        raise ApiError(400, "UNSUPPORTED_FILE_TYPE", "仅支持 CSV、XLSX 和 XLS 文件")
    raw_rows = {
        ".csv": _read_csv,
        ".xlsx": _read_xlsx,
        ".xls": _read_xls,
    }[extension](path)
    if not raw_rows:
        raise _parse_error()

    headers: dict[int, str] = {}
    for index, value in enumerate(raw_rows[0]):
        normalized = _normalize_header(value)
        if normalized is not None:
            if normalized in headers.values():
                raise _parse_error()
            headers[index] = normalized
    if not REQUIRED_HEADERS.issubset(headers.values()):
        raise _parse_error()

    rows: list[dict[str, str | int | None]] = []
    for row_number, values in enumerate(raw_rows[1:], start=2):
        normalized_values = {field: "" for field in FIELD_LIMITS}
        for index, field in headers.items():
            normalized_values[field] = _cell_text(values[index] if index < len(values) else None)
        if not any(normalized_values.values()):
            continue
        rows.append({"row_number": row_number, **normalized_values})
    if len(rows) > MAX_ROWS:
        raise ApiError(400, "ROSTER_ROW_LIMIT_EXCEEDED", "名单数据不能超过 500 行")

    seen: set[str] = set()
    validation_rows = [_validation_row(row, target_class_name, seen) for row in rows]
    return ParsedRoster(
        rows=rows,
        validation={"summary": _summarize(validation_rows), "rows": validation_rows},
    )


def summarize_rows(rows: list[dict[str, object]]) -> dict[str, int]:
    return _summarize(rows)
